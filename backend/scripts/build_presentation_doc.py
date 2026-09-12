"""One-off script to generate PRESENTATION.docx from hardcoded content
mirroring the actual codebase (prompts copied verbatim from
app/analysis/gemini_client.py as of the time this was written). Not part
of the running product -- run manually when the doc needs regenerating.
"""
from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Pt, RGBColor, Inches
from docx.enum.table import WD_TABLE_ALIGNMENT

doc = Document()

# ---- base style ----
normal = doc.styles["Normal"]
normal.font.name = "Calibri"
normal.font.size = Pt(11)

ACCENT = RGBColor(0x5B, 0x21, 0xB6)  # deep purple
DARK = RGBColor(0x1A, 0x1D, 0x2B)


def h1(text):
    p = doc.add_heading(text, level=1)
    p.runs[0].font.color.rgb = ACCENT
    return p


def h2(text):
    p = doc.add_heading(text, level=2)
    p.runs[0].font.color.rgb = ACCENT
    return p


def h3(text):
    p = doc.add_heading(text, level=3)
    p.runs[0].font.color.rgb = DARK
    return p


def para(text, bold=False, italic=False, size=11):
    p = doc.add_paragraph()
    r = p.add_run(text)
    r.bold = bold
    r.italic = italic
    r.font.size = Pt(size)
    return p


def bullets(items):
    for item in items:
        doc.add_paragraph(item, style="List Bullet")


def numbered(items):
    for item in items:
        doc.add_paragraph(item, style="List Number")


def code_block(text):
    p = doc.add_paragraph()
    p.paragraph_format.left_indent = Inches(0.3)
    for i, line in enumerate(text.split("\n")):
        r = p.add_run(("\n" if i else "") + line)
        r.font.name = "Consolas"
        r.font.size = Pt(9.5)
        r.font.color.rgb = RGBColor(0x2E, 0x2E, 0x2E)
    # light shading on the paragraph
    from docx.oxml.ns import qn
    from docx.oxml import OxmlElement
    shd = OxmlElement("w:shd")
    shd.set(qn("w:fill"), "F3F1FA")
    p._p.get_or_add_pPr().append(shd)
    return p


def rule_row(table, name, logic, threshold):
    row = table.add_row().cells
    row[0].text = name
    row[1].text = logic
    row[2].text = threshold


# =========================================================================
# TITLE PAGE
# =========================================================================
title = doc.add_paragraph()
title.alignment = WD_ALIGN_PARAGRAPH.CENTER
r = title.add_run("AI-Powered Video Retention\nDrop Analysis System")
r.font.size = Pt(28)
r.bold = True
r.font.color.rgb = ACCENT
doc.add_paragraph()

sub = doc.add_paragraph()
sub.alignment = WD_ALIGN_PARAGRAPH.CENTER
r = sub.add_run("A deterministic feature-extraction pipeline combined with Gemini-based\n"
                "visual/audio reasoning to explain WHY short-form video retention drops,\n"
                "with cited evidence and concrete suggestions.")
r.font.size = Pt(13)
r.italic = True
doc.add_paragraph()
doc.add_paragraph()

info = doc.add_paragraph()
info.alignment = WD_ALIGN_PARAGRAPH.CENTER
info.add_run("Demo / Technical Documentation\nPrepared for internal review").font.size = Pt(11)

doc.add_page_break()

# =========================================================================
# EXECUTIVE SUMMARY
# =========================================================================
h1("1. Executive Summary")
para(
    "This project analyzes a short-form video (a reel) together with its per-second viewer-retention "
    "curve, automatically identifies the specific moments where retention dropped abnormally, and "
    "explains WHY -- using a combination of classical computer-vision / audio-signal analysis and "
    "Google Gemini's multimodal (video + audio) reasoning. The output is a ranked, evidence-backed "
    "report with concrete editing suggestions, plus an interactive web UI where the reel and its "
    "retention graph play back in sync."
)
para(
    "The design goal from day one was to generalize beyond Instagram Reels specifically: the same "
    "pipeline is meant to review AI-generated video content for scenes that are not 'on par' with "
    "their neighbors (a face that renders inconsistently, a shot that looks degraded compared to the "
    "rest of the video), correlated against real audience retention data -- not just a generic quality "
    "checker."
)
h3("What makes this more than a wrapper around an LLM call")
bullets([
    "A full deterministic feature-extraction layer (image quality, shot-cut detection, audio loudness, "
    "speech-to-text, semantic dialogue-repetition detection) runs entirely locally and for free -- no "
    "API cost -- before any AI model is ever called.",
    "Gemini is only invoked for the specific few seconds around an actual detected anomaly, evaluated "
    "against real before/after video+audio clips -- never the whole video, never blindly.",
    "Every threshold and design decision below was tuned or corrected against real, observed failures "
    "during development, not guessed -- concrete examples are called out in Section 7.",
])

doc.add_page_break()

# =========================================================================
# PROBLEM STATEMENT
# =========================================================================
h1("2. Problem Statement")
para(
    "Short-form video platforms (Instagram Reels, TikTok, YouTube Shorts) show creators a per-second "
    "retention curve -- what percentage of viewers were still watching at each moment -- but never "
    "explain WHY retention fell where it did. Creators are left guessing whether a drop was caused by "
    "a boring hook, a jarring cut, an audio problem, poor visual quality, or something else entirely."
)
para("This system automates that diagnosis:")
numbered([
    "Detect exactly where retention fell faster than the video's own normal pace.",
    "Gather real evidence (actual video/audio clips, plus deterministic quality/audio/edit metrics) "
    "for each of those moments.",
    "Have an AI model watch and listen to the evidence and judge whether something looks or sounds "
    "wrong -- and if not, still report what else might explain the drop.",
    "Synthesize everything into one prioritized, specific report with suggested fixes.",
])
para(
    "A key constraint: Instagram's retention graph has no public API -- it only exists inside the "
    "Instagram Edits app UI and only reveals exact numbers via a drag/tap gesture. The pipeline "
    "extracts this by OCR-reading a screen recording of that gesture (see Section 4.1)."
)

doc.add_page_break()

# =========================================================================
# ARCHITECTURE
# =========================================================================
h1("3. System Architecture")
para("End-to-end data flow:")
code_block(
    "Retention recording --[OCR]--> RetentionCurve --+\n"
    "                                                  |\n"
    "                                                  v\n"
    "Reel URL --[yt-dlp + cookies]--> Video file --> Feature Timeline\n"
    "                                                  |  (per-second, local, free, no LLM calls:\n"
    "                                                  |   image quality, shot cuts, audio, transcript)\n"
    "                                                  v\n"
    "                                    Drop / Rise / Skip Detection\n"
    "                                    (robust statistics vs. this video's own baseline)\n"
    "                                                  |\n"
    "                                                  v\n"
    "                              Evidence Gathering (real before/during video+audio clips)\n"
    "                                                  |\n"
    "                                                  v\n"
    "                         Gemini Per-Event Analysis (visual + audio judgment)\n"
    "                                                  |\n"
    "                                                  v\n"
    "                          Gemini Synthesis Pass (text-only, ranks + writes report)\n"
    "                                                  |\n"
    "                                                  v\n"
    "                                    Interactive Web UI + report.md"
)
para(
    "Everything above the 'Gemini Per-Event Analysis' line is deterministic and runs at zero API cost. "
    "Only the handful of seconds around an actual detected drop event is ever sent to the AI model -- "
    "this keeps cost low and, more importantly, keeps the AI's job focused on judgment it's actually "
    "good at (visual/audio perception) rather than pattern-mining a whole video."
)

h2("3.1 Module Layout")
mods = [
    ("app/ingest/", "video_fetch.py (download the reel), retention_extract.py (OCR the retention graph)"),
    ("app/features/", "quality.py, shots.py, audio.py, dialogue.py, timeline.py -- the deterministic per-second feature layer"),
    ("app/analysis/", "drops.py (event detection), evidence.py (clip extraction), gemini_client.py (AI calls), run.py (orchestrates analysis)"),
    ("app/agent/", "orchestrator.py -- chains every stage into one callable pipeline"),
    ("app/web/", "FastAPI backend + interactive single-page frontend"),
]
table = doc.add_table(rows=1, cols=2)
table.style = "Light Grid Accent 5"
table.rows[0].cells[0].text = "Module"
table.rows[0].cells[1].text = "Responsibility"
for name, desc in mods:
    row = table.add_row().cells
    row[0].text = name
    row[1].text = desc

doc.add_page_break()

# =========================================================================
# TECHNICAL PIPELINE DEEP DIVE
# =========================================================================
h1("4. Technical Pipeline, Stage by Stage")

h2("4.1 Retention Curve Extraction (OCR)")
para(
    "Instagram's retention graph only exists inside the Edits app and only reveals exact numbers via "
    "a drag/tap gesture that pops up a tooltip (e.g. '62% at 0:04'). There is no API. The workflow: "
    "screen-record yourself slowly dragging across the graph, and the pipeline OCRs the tooltip out of "
    "each video frame."
)
h3("Auto-calibrating crop detection")
para(
    "Different recording setups (a raw phone screenshot vs. a screen-mirrored capture) place the graph "
    "at different positions in the frame -- a hardcoded crop region broke immediately across two real "
    "test sources. The fix locates the '100%' axis gridline label (present in every frame, at a fixed "
    "position) by VOTING across ~8 sampled frames, then derives the tooltip-search region as everything "
    "above that gridline. A single-frame check isn't reliable: at t=0 the tooltip itself also reads "
    "'100%', colliding with the axis label -- voting resolves this because the axis label is static "
    "across every frame while the tooltip's '100%' only appears once, transiently."
)
h3("Automatic outlier rejection")
bullets([
    "A timestamp reading beyond the reel's own actual duration is dropped (a digit misread cannot be "
    "a real timestamp). The reel's true duration is itself auto-detected the same voting way.",
    "A single percentage reading that both its neighbors disagree with (while agreeing with each "
    "other) is dropped as a misread digit -- but a genuine single-step change is preserved, since a "
    "real rewatch spike would not have both neighbors contradicting it.",
])

h2("4.2 Video Ingestion")
para(
    "yt-dlp handles the download. Chrome's cookie encryption (App-Bound Encryption, introduced in "
    "Chrome 127+) blocks reading its cookie store from any external tool -- confirmed during "
    "development even after installing the relevant decryption dependencies. The working authentication "
    "path is a manually exported cookies.txt file (Netscape format), stored server-side so the person "
    "using the tool never has to provide it."
)

h2("4.3 Feature Timeline (all deterministic, zero API cost)")
h3("Image quality")
para(
    "Four per-frame metrics, sampled at 5fps: sharpness (variance of Laplacian), blockiness "
    "(compression-artifact proxy -- ratio of pixel discontinuity at the 8-pixel encoder grid vs. "
    "elsewhere), high-frequency energy ratio (catches effective-resolution loss / upscaling), and "
    "brightness."
)
para(
    "Critically, each frame is scored against its OWN SHOT's mean/standard deviation, not the whole "
    "video's. Scoring against a global baseline conflates 'this frame is blurry relative to its "
    "neighbors' (a real defect) with 'this shot just looks different from the video average' (which "
    "could simply be a new location or lighting -- not a defect). This was corrected mid-project after "
    "being identified as a real gap (see Section 7)."
)
h3("Shot-cut analysis")
para(
    "PySceneDetect finds cut boundaries. CLIP (ViT-B-32, OpenAI weights) embeds each shot's middle "
    "frame for similarity comparison, producing two distinct flags: a 'redundant cut' (a cut to "
    "near-identical content -- 92%+ similarity to the immediately preceding shot) and a 'reused shot' "
    "(the same footage reappearing later -- 90%+ similarity to any non-adjacent shot)."
)
h3("Audio")
para(
    "A loudness curve (RMS-based) flags silence and abrupt volume jumps. Whisper (faster-whisper) "
    "provides speech transcription with word-level timestamps. Language is auto-detected by default "
    "(not assumed), with automatic escalation to a larger model if detection confidence is low. Every "
    "transcribed segment must clear an explicit 80%-speech-confidence bar before being trusted at all "
    "-- see Section 5 for the exact rule and the real test case that calibrated it."
)
h3("Dialogue redundancy")
para(
    "Sentence-embedding similarity (a dedicated sentence-transformer model, not CLIP's text encoder) "
    "flags when a speaker repeats a statement -- catches paraphrases, not just exact repeats."
)

h2("4.4 Drop / Rise / Skip Detection")
para(
    "Retention 'events' are detected using a robust statistical threshold (median + MAD) on the "
    "per-second rate of change -- not a fixed percentage. Every video's retention is always trending "
    "down; a fixed threshold would flag noise. Only a rate that is abnormal relative to THIS VIDEO'S "
    "OWN typical pace counts as an event."
)
para(
    "A 2-second lookback window is applied before each detected drop: a viewer decides to leave some "
    "time after whatever triggered it, not the instant retention numbers move, so evidence gathering "
    "reaches slightly earlier than the statistically-detected drop point."
)
para(
    "The same detection logic also runs in reverse to find RISES (retention recovering) -- relevant on "
    "platforms/players where retention is built from playhead-position samples: a viewer who skips "
    "ahead stops being counted for the skipped seconds and resumes being counted at the landing point, "
    "producing a dip-then-partial-recovery shape. A drop closely followed by a substantial recovery is "
    "reclassified as a likely SKIP rather than genuine disengagement, so the report doesn't overstate "
    "its severity."
)

h2("4.5 Evidence Gathering")
para(
    "For each detected event, real short video clips (with audio) are extracted -- a 'BEFORE' clip "
    "(a clean reference period, prior to the investigation window) and a 'DURING' clip (the drop "
    "window itself). This was an explicit upgrade over an earlier still-frame-based approach: static "
    "frames cannot show motion, temporal continuity, or audio-visual sync -- exactly the properties "
    "that matter most for catching AI-generation artifacts (a face that morphs across frames, "
    "flickering, geometry that shifts within one shot)."
)

doc.add_page_break()

# =========================================================================
# RULES BEHIND THE SUGGESTIONS
# =========================================================================
h1("5. Rules Behind the Suggestions")
para(
    "Every flag and suggestion the system produces traces back to one of the following explicit, "
    "tunable rules. None of these thresholds were guessed blind -- each was set or corrected against "
    "a concrete observed case during development."
)

table = doc.add_table(rows=1, cols=3)
table.style = "Light Grid Accent 5"
hdr = table.rows[0].cells
hdr[0].text = "Rule"
hdr[1].text = "Logic"
hdr[2].text = "Threshold"

rule_row(table, "Drop event detection",
         "Per-second retention loss rate, z-scored against this video's own median/MAD (robust to outliers)",
         "z > 1.5")
rule_row(table, "Rise / recovery detection",
         "Same statistic, opposite tail -- retention increasing faster than this video's typical pace",
         "z < -1.5")
rule_row(table, "Skip vs. disengagement",
         "A drop followed by a rise within a short window that recovers most of what was lost is "
         "reclassified as a likely skip, not full disengagement",
         "rise within 20s, recovers >=50% of pct lost")
rule_row(table, "Investigation lag window",
         "Evidence is gathered starting before the statistically-detected drop point (viewers react "
         "with a delay), through slightly after it ends",
         "-2.0s before, +0.5s after")
rule_row(table, "Per-shot quality anomaly",
         "A frame's sharpness/blockiness is compared to its OWN shot's mean/std, not the whole video's "
         "-- isolates real mid-shot glitches from ordinary scene-to-scene variation",
         "z > 1.5 (blockiness), z < -1.5 (sharpness)")
rule_row(table, "Redundant cut",
         "CLIP visual-similarity of a shot vs. the immediately preceding shot",
         ">= 0.92 cosine similarity")
rule_row(table, "Reused shot",
         "CLIP visual-similarity of a shot vs. any OTHER non-adjacent shot elsewhere in the video",
         ">= 0.90 cosine similarity")
rule_row(table, "Redundant dialogue",
         "Sentence-embedding semantic similarity between two spoken segments (min. 3 words each, to "
         "avoid flagging short filler phrases)",
         ">= 0.82 cosine similarity")
rule_row(table, "Trusted speech/lyrics",
         "A transcribed segment is only kept if the model is confident it is actually speech, AND "
         "confident in the specific words -- otherwise treated as 'no reliable transcript', never guessed",
         "speech confidence > 80%, avg. token log-probability > -1.0")
rule_row(table, "AI visual/audio quality judgment",
         "Gemini watches the BEFORE/DURING clips together (motion + audio, not stills) and judges "
         "whether something looks or sounds genuinely wrong -- a content change alone is explicitly "
         "NOT sufficient to trigger this",
         "model judgment, structured output (see Section 6)")
rule_row(table, "Investigation budget",
         "Only the top-N drop events by actual percentage-points-lost are sent for deep AI analysis, "
         "to control cost -- all events are still detected and listed",
         "default N = 3")

para("")
para(
    "Why the 80%-speech-confidence rule matters as a concrete example: during testing, a background "
    "song's vocals were suspected to contain audible lyrics. Forcing the transcription model past its "
    "own safety threshold produced text that was technically valid script but semantically incoherent, "
    "with the model itself reporting only 9-17% confidence that it was even hearing speech. Using a "
    "larger model on an isolated vocal track (via audio source separation) improved confidence "
    "substantially, but still landed at only 26-41% -- still correctly rejected under the 80% bar. "
    "The system is deliberately built to say 'no reliable transcript' rather than guess and risk being "
    "wrong in front of a client."
)

doc.add_page_break()

# =========================================================================
# EXACT GEMINI PROMPTS
# =========================================================================
h1("6. AI Analysis: Exact Prompts Used")
para(
    "The system uses Google Gemini (gemini-flash-latest) in two distinct passes with two different "
    "prompts. Both request STRUCTURED (schema-validated) JSON output, not free text, so results are "
    "always machine-parseable. The literal prompt templates are reproduced below, exactly as they "
    "exist in the codebase (app/analysis/gemini_client.py) at the time of writing; {curly braces} mark "
    "where real data is substituted in at run time."
)

h2("6.1 Pass 1 -- Per-Event Visual & Audio Analysis")
para(
    "Run once per detected drop event. Two real video clips (with audio) are uploaded to Gemini's "
    "Files API and included directly in the request alongside this prompt.", italic=True
)
code_block(
'''You are reviewing a short video (a reel) to understand why viewer retention dropped sharply at one point.

You are given two short video clips, WITH AUDIO, in order:
- BEFORE: t={before_start}-{before_end}s, just before the drop window
- DURING: t={during_start}-{during_end}s, the drop window itself

Retention went from {retention_before}% to {retention_after}% of viewers still watching across this window
({pct_points_lost} percentage points lost) -- much steeper than this video's own typical pace
(z={peak_drop_z}, where >1.5 is already unusual for this video).

YOUR TASK: watch both clips -- motion and audio together, not just individual frames -- and compare
DURING against BEFORE as a viewer would. Judge whether DURING shows a genuine quality or rendering
problem: reduced clarity/sharpness, blurred or malformed characters/faces/hands/objects, flicker or
morphing across frames, broken temporal continuity, warped geometry, texture breakdown, audio that
doesn't sync to what's on screen, or anything that looks/sounds like it wasn't shot/rendered/generated
properly. Do NOT assume the content is real footage or AI-generated -- judge only on what you actually
see and hear. A content/subject change alone (new location, new shot, new topic) is NOT a quality issue
by itself; only flag it if something about HOW it's rendered or played looks or sounds wrong.

Supporting numeric signals (context only -- your own read of the clips is what matters most):
{ numeric evidence: quality metrics, audio levels, cut flags, transcript -- as JSON }

Respond with the structured fields requested.'''
)
h3("Structured output schema (Pass 1)")
bullets([
    "quality_issue_found (boolean)",
    "description (plain-language explanation of what looks wrong)",
    "affected_elements (list -- e.g. 'character face', 'hands', 'lighting continuity')",
    "severity (none / low / medium / high)",
    "confidence (0-1)",
    "likely_explains_retention_drop (boolean)",
    "reasoning (how the evidence connects to the drop)",
    "other_observations (anything else notable, even if NOT a quality issue -- e.g. a tone or audio "
    "shift; this field is what surfaces real causes when the quality check itself comes back negative)",
    "suggestion (concrete fix)",
])

h2("6.2 Pass 2 -- Synthesis (text-only, no video)")
para(
    "Run once per video, after all per-event analyses complete. Receives the full per-second feature "
    "digest for the ENTIRE video (not just the investigated events) plus every Pass-1 finding as JSON.",
    italic=True
)
code_block(
'''You are writing the final analysis report for a short video (a reel), explaining why viewer
retention dropped where it did and what to do about it. You're given:

1. A per-second digest of the whole reel (retention, cuts, quality flags, audio, transcript):
---
{ full per-second timeline digest }
---

2. Structured findings already produced for each major drop event (a vision model already compared
the actual video frames before/during each drop and recorded what it saw -- treat "other_observations"
in each as real evidence, not speculation, since it came from looking at the actual footage):
---
{ Pass-1 JSON findings for every investigated event }
---

YOUR TASK: synthesize this into one coherent, prioritized report. Rank the drop events by how many
percentage points of retention they actually cost (biggest impact first, not necessarily chronological
order). For each, give a concrete, specific reason grounded in the evidence above -- not generic advice.
Then give a small number of cross-cutting suggestions that address the reel as a whole, not just one
moment.

IMPORTANT: if an event's JSON has a non-null "likely_skip" field, retention recovered shortly
afterward -- this is consistent with viewers skipping ahead rather than genuinely disengaging. Describe
it as a likely skip, not full disengagement, and don't overstate its severity as if the audience left
for good, but still suggest why that segment specifically might have invited a skip (it's still worth
improving).

Write for someone who made this reel and wants to know exactly what to change next time. Be direct and
specific (reference actual timestamps and numbers), not vague ("improve pacing" is not useful; "the
opening text card holds for 3+ seconds with no motion or audio before the first cut -- cut that to
under 1.5s" is).'''
)
h3("Structured output schema (Pass 2)")
bullets([
    "headline (one sentence: the single biggest reason this reel lost viewers)",
    "overview (2-4 sentence narrative of the whole retention story)",
    "ranked_findings (list, each with: rank, window, pct_points_lost, cause_summary, evidence, suggestion)",
    "overall_suggestions (cross-cutting recommendations beyond individual events)",
])

h2("6.3 Error-Handling Policy")
para(
    "Two distinct Gemini API failure modes are handled differently, deliberately: a 503 ('high demand') "
    "is transient and is retried with exponential backoff (up to 6 attempts: 3s, 6s, 12s, 24s, 48s). A "
    "429 (quota exhausted) is NOT transient -- retrying only burns more requests against an "
    "already-exhausted budget, so the system fails fast with an actionable message instead. This "
    "distinction was added after observing that blind retries on a 429 made an outage longer, not "
    "shorter, during testing."
)

doc.add_page_break()

# =========================================================================
# TECH STACK
# =========================================================================
h1("7. Technology Stack & Engineering Notes")
h2("7.1 Stack")
table = doc.add_table(rows=1, cols=2)
table.style = "Light Grid Accent 5"
table.rows[0].cells[0].text = "Layer"
table.rows[0].cells[1].text = "Technology"
stack = [
    ("Video download", "yt-dlp + manually-exported cookies.txt (Chrome's App-Bound Encryption blocks --cookies-from-browser)"),
    ("Retention OCR", "OpenCV + Tesseract OCR, auto-calibrating crop/duration detection"),
    ("Image quality", "OpenCV (Laplacian variance, FFT analysis, custom blockiness metric)"),
    ("Shot detection & similarity", "PySceneDetect + OpenCLIP (ViT-B-32, OpenAI weights)"),
    ("Speech-to-text", "faster-whisper (small/medium, auto language detection + confidence-gated escalation)"),
    ("Vocal isolation", "Demucs (htdemucs) -- source separation for lyrics under music"),
    ("Dialogue similarity", "sentence-transformers (all-MiniLM-L6-v2)"),
    ("AI visual/audio/text reasoning", "Google Gemini (gemini-flash-latest) via google-genai SDK, structured JSON output"),
    ("Backend", "FastAPI (Python), background-threaded job processing"),
    ("Frontend", "Vanilla HTML/CSS/JS, hand-built SVG chart -- no framework/build step"),
    ("Tunneling / sharing", "Cloudflare Tunnel (cloudflared quick tunnel, no account needed)"),
]
for k, v in stack:
    row = table.add_row().cells
    row[0].text = k
    row[1].text = v

h2("7.2 Notable Problems Found and Fixed During Development")
para(
    "These are concrete examples of debugging/engineering rigor applied while building this, not "
    "just a list of features:", italic=True
)
bullets([
    "Global vs. per-shot quality scoring: an early version flagged quality issues by comparing every "
    "frame to the whole video's average. Identified as a real design flaw (a legitimately different "
    "scene would be falsely flagged) and corrected to score each frame against its own shot's "
    "baseline.",
    "Chrome cookie decryption: --cookies-from-browser chrome failed with a DPAPI decryption error. "
    "Traced to Chrome 127+'s App-Bound Encryption (ties decryption to Chrome's own elevation service); "
    "confirmed unfixable externally even after installing the relevant crypto dependencies, and "
    "switched to a manually-exported cookies.txt.",
    "Whisper model caching bug: a global cache stored one model instance regardless of requested size, "
    "which would have silently defeated automatic escalation to a larger model on low-confidence "
    "results. Found and fixed before it could cause a silent failure.",
    "Skip-detection gap threshold: an initial 5-second maximum gap between a drop and its recovery was "
    "too tight for the target scenario (a 10-second skip) -- caught by testing against synthetic data "
    "modeling that exact case, and corrected to 20 seconds.",
    "429 vs. 503 API errors: initially retried both the same way. Real testing showed blind-retrying a "
    "quota-exhausted (429) error only worsened the situation; separated the handling so quota errors "
    "fail fast instead.",
    "Security review before public sharing: before exposing the tool via a public tunnel, an "
    "unauthenticated file-serving endpoint was found to be capable of returning the raw uploaded "
    "cookies.txt (a live session token) to anyone who knew a job ID. Fixed to explicitly refuse "
    "serving back any uploaded input, only pipeline-produced outputs.",
])

doc.add_page_break()

# =========================================================================
# LIMITATIONS
# =========================================================================
h1("8. Known Limitations & Future Work")
bullets([
    "Only one AI-reasoned rule is active today (generic visual/audio quality vs. neighboring content). "
    "The deterministic redundant-cut / reused-shot / redundant-dialogue detectors already run and "
    "appear in the internal data, but are not yet each individually escalated to Gemini for a full "
    "visual explanation -- the same before/during-evidence pattern would extend directly to them.",
    "Reliable song-lyric transcription needs audio-fingerprint + lyrics-database lookup (the track is "
    "usually an identifiable existing song), which is a stronger approach than transcription-based ASR "
    "and is a recommended next addition.",
    "The current 'agent' is a deterministic sequencer (fetch -> extract -> analyze -> report) with no "
    "standalone CLI entry point yet -- it does not autonomously watch an account for new content; that "
    "would be a separate, larger addition requiring platform-side polling.",
    "The web backend is single-process with in-memory + disk-recoverable job tracking, appropriate for "
    "a demo/single-operator tool, not yet multi-tenant production infrastructure.",
    "No authentication layer exists in the web UI today.",
])

doc.add_page_break()

# =========================================================================
# HOW TO RUN THE DEMO
# =========================================================================
h1("9. Running the Live Demo")

h2("9.1 Start the Website")
para("From the project's backend directory, in PowerShell:", italic=True)
code_block(
    "cd backend\n"
    ".venv\\Scripts\\python -m uvicorn app.web.main:app --host 127.0.0.1 --port 8420"
)
para("Then open in a browser:")
code_block("http://127.0.0.1:8420")
para(
    "Leave this terminal window open -- the server runs in it. A visitor pastes a reel URL, uploads a "
    "retention-graph recording (a sample and recording instructions are shown directly in the page), "
    "and clicks Run Analysis. Instagram credentials are handled entirely server-side; nothing is asked "
    "of the visitor."
)

h2("9.2 Stop the Website")
para("In the terminal running the server, press:")
code_block("Ctrl + C")
para("Or, to force-stop from another terminal if needed:")
code_block(
    "tasklist | findstr python\n"
    "taskkill /F /PID <the process ID shown>"
)

h2("9.3 Create a Public Tunnel (to share with someone remote)")
para("A standalone copy of cloudflared is kept at backend\\.bin\\cloudflared.exe (no install needed). "
     "With the website already running, in a NEW terminal window:")
code_block(
    "cd backend\n"
    ".bin\\cloudflared.exe tunnel --url http://127.0.0.1:8420"
)
para(
    "This prints a public HTTPS URL in the form https://<random-words>.trycloudflare.com -- it stays "
    "live only as long as this terminal window is open."
)

h2("9.4 Stop the Tunnel")
para("In the terminal running cloudflared, press:")
code_block("Ctrl + C")
para("The public URL stops working immediately; the local website keeps running unaffected.")

h2("9.5 Sharing the Link")
para(
    "There are two ways to share, and they behave differently -- worth choosing deliberately:"
)
table = doc.add_table(rows=1, cols=2)
table.style = "Light Grid Accent 5"
table.rows[0].cells[0].text = "Link form"
table.rows[0].cells[1].text = "Behavior"
row = table.add_row().cells
row[0].text = "https://<tunnel-url>/"
row[1].text = "Full form -- the recipient can submit their OWN reel + recording for a fresh analysis. Uses the shared Gemini API quota and the server-side Instagram session for every request."
row = table.add_row().cells
row[0].text = "https://<tunnel-url>/?job=<job_id>"
row[1].text = "View-only -- opens directly into an already-completed result (video + synced chart + report). No new API calls, no new Instagram fetch. Recommended when the goal is just to SHOW a finished analysis."
para("")
para(
    "Security note: the tunnel has no login. Anyone with either link can use it. Do not post the bare "
    "form URL publicly -- only share it with people you trust, since it consumes your own API quota "
    "and Instagram session on their behalf.", italic=True
)

doc.save(r"C:\Users\rithi\Desktop\video_retention\PRESENTATION.docx")
print("Saved PRESENTATION.docx")
